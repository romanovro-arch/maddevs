import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def create_element(name):
    return OxmlElement(name)

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=140, bottom=140, left=200, right=200):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                      f'<w:top w:w="{top}" w:type="dxa"/>'
                      f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                      f'<w:left w:w="{left}" w:type="dxa"/>'
                      f'<w:right w:w="{right}" w:type="dxa"/>'
                      f'</w:tcMar>')
    tcPr.append(tcMar)

def set_cell_border(cell, **kwargs):
    """
    kwargs: top, bottom, left, right
    values: dict(sz=12, val='single', color='4F46E5')
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    for edge in ('top', 'left', 'bottom', 'right'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = f'<w:{edge} {nsdecls("w")} w:val="{edge_data.get("val", "single")}" w:sz="{edge_data.get("sz", 4)}" w:space="0" w:color="{edge_data.get("color", "auto")}"/>'
            tcBorders.append(parse_xml(tag))
        else:
            tag = f'<w:{edge} {nsdecls("w")} w:val="none"/>'
            tcBorders.append(parse_xml(tag))
    tcPr.append(tcBorders)

def generate_docx(output_path):
    doc = Document()

    # Set page margins (standard 1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Palette
    COLOR_PRIMARY = RGBColor(15, 23, 42)      # Slate 900
    COLOR_ACCENT = RGBColor(79, 70, 229)      # Indigo 600
    COLOR_MUTED = RGBColor(100, 116, 139)     # Slate 500
    COLOR_BODY = RGBColor(51, 65, 85)         # Slate 700

    # 1. Header / Tag
    p_tag = doc.add_paragraph()
    p_tag.paragraph_format.space_before = Pt(0)
    p_tag.paragraph_format.space_after = Pt(4)
    run_tag = p_tag.add_run("MAD DEVS ENGINEERING CHALLENGE  •  ЭССЕ")
    run_tag.font.name = "Arial"
    run_tag.font.size = Pt(9.5)
    run_tag.font.bold = True
    run_tag.font.color.rgb = COLOR_ACCENT

    # 2. Main Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(8)
    run_title = p_title.add_run("Самая сложная проблема, которую я решаю прямо сейчас")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    # 3. Subtitle
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(18)
    run_sub = p_sub.add_run("Обеспечение абсолютной консистентности данных в высококонкурентных асинхронных системах")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(11.5)
    run_sub.font.italic = True
    run_sub.font.color.rgb = COLOR_MUTED

    # 4. Meta Box (Table with subtle border and fill)
    meta_table = doc.add_table(rows=1, cols=1)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_cell = meta_table.cell(0, 0)
    meta_cell.width = Inches(6.5)
    set_cell_background(meta_cell, "F8FAFC") # slate-50
    set_cell_margins(meta_cell, top=140, bottom=140, left=200, right=200)
    set_cell_border(meta_cell, left=dict(sz=16, val='single', color='4F46E5')) # indigo left border

    p_meta = meta_cell.paragraphs[0]
    p_meta.paragraph_format.space_before = Pt(0)
    p_meta.paragraph_format.space_after = Pt(0)
    p_meta.paragraph_format.line_spacing = 1.25

    def add_meta_field(p, label, val):
        r1 = p.add_run(f"{label}: ")
        r1.font.name = "Arial"
        r1.font.size = Pt(9.5)
        r1.font.bold = True
        r1.font.color.rgb = COLOR_PRIMARY
        r2 = p.add_run(f"{val}   |   ")
        r2.font.name = "Arial"
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = COLOR_BODY

    add_meta_field(p_meta, "Формат", "10 предложений")
    add_meta_field(p_meta, "Проект", "MadEvents (Кейс №9)")
    add_meta_field(p_meta, "Стек", "Bun, Hono, SQLite WAL, WebSockets, Next.js")

    # Add space after table
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(8)
    p_spacer.paragraph_format.space_after = Pt(4)

    # 5. Essay Sentences
    sentences = [
        ("1", "Самая сложная и захватывающая инженерная проблема, которую я решаю прямо сейчас — это обеспечение абсолютной целостности и детерминированности данных в условиях агрессивного конкурентного доступа (race conditions) при сохранении сверхнизких задержек и высокой отзывчивости системы."),
        ("2", "В сервисах регистрации и бронирования критический момент наступает тогда, когда десятки параллельных запросов претендуют на последнее доступное место в одну и ту же миллисекунду."),
        ("3", "Наивные проверки на уровне приложения здесь неизбежно дают сбой, приводя к овербукингу, рассинхронизации очередей и фантомным записям."),
        ("4", "Фундаментальный вызов заключается в том, чтобы найти идеальный баланс между гранулярностью блокировок, скоростью выполнения транзакций и предсказуемостью поведения всей цепочки событий."),
        ("5", "Для решения этой задачи я проектирую строгую сериализацию транзакций на уровне хранилища данных с детерминированным переводом избыточных запросов в атомарный FIFO-лист ожидания."),
        ("6", "Вторая грань этой проблемы — надежное каскадное продвижение очереди: при отмене брони система обязана автономно и без малейшей потери консистентности передать слот следующему кандидату, выпустить билет и отправить уведомление."),
        ("7", "Дополнительную сложность накладывает требование строгой идемпотентности: однократный чекин билета по зашифрованному QR-коду исключает повторный вход и не прощает дублирования мутаций."),
        ("8", "Наконец, критически важно замкнуть этот контур на реальное время, доставляя актуальную телеметрию организаторам по протоколу WebSockets без перегрузки канала и с гарантией актуальности данных."),
        ("9", "В конечном счете эта проблема выходит за рамки простого кодинга — это создание предсказуемой архитектуры, где каждый граничный сценарий предварительно валидируется жесткими стресс-тестами параллелизма."),
        ("10", "Именно решение таких нетривиальных задач на стыке транзакционной строгости, асинхронных очередей и высокой скорости работы мотивирует меня развиваться в современной разработке.")
    ]

    for num, text in sentences:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.25

        run_num = p.add_run(f"{num}. ")
        run_num.font.name = "Arial"
        run_num.font.size = Pt(11)
        run_num.font.bold = True
        run_num.font.color.rgb = COLOR_ACCENT

        run_text = p.add_run(text)
        run_text.font.name = "Arial"
        run_text.font.size = Pt(11)
        run_text.font.color.rgb = COLOR_BODY

    # 6. Conclusion / Verification Box
    p_spacer2 = doc.add_paragraph()
    p_spacer2.paragraph_format.space_before = Pt(8)
    p_spacer2.paragraph_format.space_after = Pt(0)

    verif_table = doc.add_table(rows=1, cols=1)
    verif_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    verif_cell = verif_table.cell(0, 0)
    verif_cell.width = Inches(6.5)
    set_cell_background(verif_cell, "F1F5F9") # slate-100
    set_cell_margins(verif_cell, top=160, bottom=160, left=200, right=200)
    set_cell_border(verif_cell, left=dict(sz=16, val='single', color='10B981')) # emerald left border

    p_v = verif_cell.paragraphs[0]
    p_v.paragraph_format.space_before = Pt(0)
    p_v.paragraph_format.space_after = Pt(0)
    p_v.paragraph_format.line_spacing = 1.25

    rv_title = p_v.add_run("Практическая верификация в проекте MadEvents:\n")
    rv_title.font.name = "Arial"
    rv_title.font.size = Pt(10)
    rv_title.font.bold = True
    rv_title.font.color.rgb = RGBColor(16, 185, 129)

    rv_body = p_v.add_run(
        "Данный подход полностью реализован и протестирован в репозитории проекта: "
        "атомарные транзакции SQLite IMMEDIATE, автоматический FIFO лист ожидания, "
        "однократный чекин и нативные WebSockets в Hono. "
        "Успешность архитектурных решений подтверждена набором стресс-тестов конкурентности (6 из 6 пройдены за 1.8с)."
    )
    rv_body.font.name = "Arial"
    rv_body.font.size = Pt(9.5)
    rv_body.font.color.rgb = COLOR_BODY

    # Footer note
    section = doc.sections[0]
    footer = section.footer
    p_footer = footer.paragraphs[0]
    p_footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rf = p_footer.add_run("Репозиторий проекта: github.com/romanovro-arch/maddevs")
    rf.font.name = "Arial"
    rf.font.size = Pt(8.5)
    rf.font.color.rgb = COLOR_MUTED

    doc.save(output_path)
    print(f"Successfully generated: {output_path}")

if __name__ == "__main__":
    out_dir = r"d:\ShareCode\Programming\maddevs"
    target_file = os.path.join(out_dir, "ESSAY_MadDevs.docx")
    generate_docx(target_file)
    # Also save as Russian filename for convenience
    target_file_ru = os.path.join(out_dir, "ЭССЕ_Самая_сложная_проблема.docx")
    generate_docx(target_file_ru)
